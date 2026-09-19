"""Inspection report generation (PDF + DOCX) from stored inspection data.

Reports are built exclusively from real database records. Missing declarations
are rendered as "Not detected"; nothing is invented or substituted.
"""

import io
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from app.config import settings
from app.models import Inspection, InspectionStatus

SECTION_TITLES = {
    "product_identity": "Product Identity",
    "manufacturer_and_marketer": "Manufacturer & Marketer",
    "dates_and_batch": "Dates & Batch",
    "pricing_and_quantity": "Pricing & Quantity",
    "ingredients_and_allergens": "Ingredients & Allergens",
    "nutritional_information": "Nutritional Information",
    "regulatory_and_certifications": "Regulatory & Certifications",
    "consumer_care": "Consumer Care",
}

NOT_DETECTED_LABEL = "Not detected"
STATUS_LABELS = {"DETECTED": "Detected", "NOT_DETECTED": "Not detected", "UNCERTAIN": "Uncertain"}
CHECK_LABELS = {"PASSED": "PASS", "FAILED": "FAIL", "REVIEW": "REVIEW", "NOT_APPLICABLE": "Not applicable"}

FINAL_VERDICT = {
    InspectionStatus.DRAFT: "Not evaluated",
    InspectionStatus.PROCESSING: "Evaluation in progress",
    InspectionStatus.FAILED: "Evaluation failed",
}
FINAL_VERDICT_CHECKS = {
    "FAILED": "NON-COMPLIANT — mandatory failures recorded",
    "REVIEW": "REVIEW REQUIRED — manual verification needed",
    "PASSED": "COMPLIANT — all evaluated parameters passed",
    "NOT_APPLICABLE": "No applicable parameters evaluated",
}


def _verdict(checks: list) -> str:
    statuses = [c.status.value for c in checks]
    for s in ("FAILED", "REVIEW", "PASSED", "NOT_APPLICABLE"):
        if s in statuses:
            return FINAL_VERDICT_CHECKS[s]
    return "No compliance checks recorded"


def _declaration_rows(inspection: Inspection) -> list[tuple[str, str, str]]:
    rows = []
    for d in sorted(inspection.declarations, key=lambda x: x.field_name):
        section = SECTION_TITLES.get(d.field_name.split(".")[0], d.field_name.split(".")[0].replace("_", " ").title())
        field = d.field_name.split(".", 1)[1].replace("_", " ").title()
        value = d.value if d.value is not None else NOT_DETECTED_LABEL
        rows.append((f"{section} — {field}", value, STATUS_LABELS.get(d.status.value, d.status.value)))
    return rows


def _final_verdict_of(inspection: Inspection) -> str:
    if inspection.status != InspectionStatus.COMPLETED:
        return FINAL_VERDICT.get(inspection.status, "Not evaluated")
    return _verdict(inspection.compliance_checks)


def _report_meta(inspection: Inspection) -> dict:
    return {
        "id": inspection.id,
        "status": inspection.status.value,
        "notes": inspection.notes,
        "inspector_id": inspection.inspector_id,
        "created_at": inspection.created_at,
        "generated_at": datetime.now(timezone.utc),
        "images": len(inspection.images),
        "violations": len(inspection.violations),
    }


def _build_pdf(inspection: Inspection) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    navy = colors.HexColor("#0f2a52")
    saffron = colors.HexColor("#e87722")
    green = colors.HexColor("#1a7a3c")
    light = colors.HexColor("#eef1f6")
    border = colors.HexColor("#c9d1de")

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Title"], fontSize=20, textColor=navy, spaceAfter=2)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, textColor=navy, spaceBefore=14, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.5, leading=11)
    verdict_style = ParagraphStyle("Verdict", parent=body, fontSize=11, textColor=navy)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title="NIRIKSH Inspection Report")
    story = []

    meta = _report_meta(inspection)
    story.append(Paragraph("NIRIKSH — Legal Metrology Compliance Inspection Report", h1))
    story.append(Paragraph(
        f"Inspection #{meta['id']:04d} &nbsp;|&nbsp; Generated: {meta['generated_at'].strftime('%d %b %Y %H:%M UTC')}",
        ParagraphStyle("sub", parent=body, textColor=colors.HexColor("#5b6472"))))

    rows = [
        ["Inspection ID", f"#{meta['id']:04d}", "Status", meta["status"]],
        ["Notes", meta["notes"] or "—", "Images captured", str(meta["images"])],
        ["Inspected on", meta["created_at"].strftime("%d %b %Y %H:%M UTC") if meta["created_at"] else "—",
         "Violations recorded", str(meta["violations"])],
    ]
    t = Table(rows, colWidths=[30 * mm, 60 * mm, 35 * mm, 47 * mm], hAlign="CENTER")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), light), ("BACKGROUND", (2, 0), (2, -1), light),
        ("GRID", (0, 0), (-1, -1), 0.5, border),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t)

    story.append(Paragraph("Final Compliance Result", h2))
    story.append(Paragraph(_final_verdict_of(inspection), verdict_style))

    if inspection.images:
        story.append(Paragraph("Evidence — Captured Product Images", h2))
        img_row = []
        for img in inspection.images[:4]:
            p = Path(img.file_path)
            if p.exists():
                img_row.append(Image(str(p), width=38 * mm, height=28 * mm))
            else:
                img_row.append(Paragraph("(image unavailable)", cell))
        while len(img_row) < 4:
            img_row.append("")
        ti = Table([img_row], colWidths=[43 * mm] * 4, hAlign="CENTER")
        ti.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, border)]))
        story.append(ti)

    decls = _declaration_rows(inspection)
    if decls:
        story.append(Paragraph("Extracted Declarations", h2))
        data = [["Declaration", "Extracted value", "Status"]]
        for name, value, status in decls:
            data.append([Paragraph(name, cell), Paragraph(str(value)[:180], cell), status])
        td = Table(data, colWidths=[70 * mm, 82 * mm, 20 * mm], hAlign="CENTER", repeatRows=1)
        td.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(td)
    else:
        story.append(Paragraph("Extracted Declarations", h2))
        story.append(Paragraph("No declarations extracted for this inspection.", body))

    if inspection.compliance_checks:
        story.append(Paragraph("Compliance Checks", h2))
        data = [["Parameter", "Result", "Findings"]]
        for c in inspection.compliance_checks:
            data.append([
                Paragraph((c.parameter or "—").replace("_", " ").upper(), cell),
                CHECK_LABELS.get(c.status.value, c.status.value),
                Paragraph(c.findings or "", cell),
            ])
        tc = Table(data, colWidths=[34 * mm, 22 * mm, 116 * mm], hAlign="CENTER", repeatRows=1)
        tc.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(tc)

    if inspection.violations:
        story.append(Paragraph("Violations", h2))
        for v in inspection.violations:
            line = f"{v.severity or '—'} · {v.error_code or '—'} — {v.description or ''}"
            story.append(KeepTogether([
                Paragraph(line, ParagraphStyle("v", parent=body, textColor=colors.HexColor("#8a2b2b"))),
                Paragraph(f"Remediation: {v.remediation}" if v.remediation else "", cell),
                Spacer(1, 4),
            ]))

    story.append(PageBreak() if len(story) > 18 else Spacer(1, 10))
    story.append(Paragraph(
        "Generated by NIRIKSH from recorded inspection data. Values shown as "
        f"\"{NOT_DETECTED_LABEL}\" were not detected and were not substituted.",
        ParagraphStyle("foot", parent=cell, textColor=colors.HexColor("#5b6472"))))

    doc.build(story)
    return buf.getvalue()


def _build_docx(inspection: Inspection) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    navy = RGBColor(0x0F, 0x2A, 0x52)
    gray = RGBColor(0x5B, 0x64, 0x72)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.name = "Calibri"

    meta = _report_meta(inspection)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NIRIKSH — Legal Metrology Compliance Inspection Report")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = navy
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"Inspection #{meta['id']:04d} | Generated: {meta['generated_at'].strftime('%d %b %Y %H:%M UTC')}")
    r.font.size = Pt(9)
    r.font.color.rgb = gray
    # editable marker for officials
    doc.add_paragraph()

    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    for label, value in [
        ("Inspection ID", f"#{meta['id']:04d}"), ("Status", meta["status"]),
        ("Notes", meta["notes"] or "—"), ("Images captured", str(meta["images"])),
        ("Inspected on", meta["created_at"].strftime("%d %b %Y %H:%M UTC") if meta["created_at"] else "—"),
        ("Violations recorded", str(meta["violations"])),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = label, value, "", ""
        for cell in (row[0], row[2]):
            for p in cell.paragraphs:
                for r2 in p.runs:
                    r2.bold = True

    doc.add_heading("Final Compliance Result", level=1)
    doc.add_paragraph(_final_verdict_of(inspection))

    if inspection.images:
        doc.add_heading("Evidence — Captured Product Images", level=1)
        for img in inspection.images[:6]:
            p = Path(img.file_path)
            if p.exists():
                doc.add_picture(str(p), width=Mm(70))

    decls = _declaration_rows(inspection)
    doc.add_heading("Extracted Declarations", level=1)
    if decls:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        for i, h in enumerate(("Declaration", "Extracted value", "Status")):
            hdr[i].text = h
            for p in hdr[i].paragraphs:
                for r2 in p.runs:
                    r2.bold = True
        for name, value, status in decls:
            row = t.add_row().cells
            row[0].text, row[1].text, row[2].text = name, str(value)[:180], status
    else:
        doc.add_paragraph("No declarations extracted for this inspection.")

    if inspection.compliance_checks:
        doc.add_heading("Compliance Checks", level=1)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        for i, h in enumerate(("Parameter", "Result", "Findings")):
            hdr[i].text = h
            for p in hdr[i].paragraphs:
                for r2 in p.runs:
                    r2.bold = True
        for c in inspection.compliance_checks:
            row = t.add_row().cells
            row[0].text = (c.parameter or "—").replace("_", " ").upper()
            row[1].text = CHECK_LABELS.get(c.status.value, c.status.value)
            row[2].text = c.findings or ""
            if c.details and c.details.get("section_reference"):
                ref = row[0].add_paragraph()
                rr = ref.add_run(f"Rule: {c.details['section_reference']}")
                rr.font.size = Pt(8)
                rr.font.color.rgb = gray

    if inspection.violations:
        doc.add_heading("Violations", level=1)
        for v in inspection.violations:
            p = doc.add_paragraph()
            r2 = p.add_run(f"{v.severity or '—'} · {v.error_code or '—'}")
            r2.bold = True
            doc.add_paragraph(v.description or "")
            if v.remediation:
                doc.add_paragraph(f"Remediation: {v.remediation}")

    doc.add_heading("Declaration Notes", level=1)
    doc.add_paragraph(
        f"Values shown as \"{NOT_DETECTED_LABEL}\" were not detected on the packaging evidence and were "
        "not substituted. Entries marked \"Uncertain\" could not be verified against raw OCR text and "
        "require manual confirmation against the source image.")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _report_path(inspection_id: int, fmt: str, content: bytes) -> Path:
    directory = settings.storage_dir / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"inspection_{inspection_id}_{fmt}.{fmt}"
    path.write_bytes(content)
    return path


def generate_report(db: Session, inspection: Inspection, fmt: str) -> Path:
    if fmt == "pdf":
        content = _build_pdf(inspection)
    elif fmt == "docx":
        content = _build_docx(inspection)
    else:
        raise ValueError(f"Unsupported report format: {fmt}")
    return _report_path(inspection.id, fmt, content)
