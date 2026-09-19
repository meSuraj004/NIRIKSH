"""Report generation tests: PDF + DOCX built from real stored inspection data."""

import io
import os
import sys
from unittest.mock import patch

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from conftest import TestSession, client, signup_and_login  # noqa: E402
from app.models import Declaration, Inspection  # noqa: E402

HEADERS = None


def setup_module(module):
    global HEADERS
    HEADERS = signup_and_login("reports@example.com")


TRANSCRIPT = """ACME JUICE MANGO
MRP Rs.65.00 (incl. of all taxes)
Net Volume 1 L
MFG 04/2026
USE BY 10/2026"""

AGGREGATED = {
    "product_identity": {"brand_name": "ACME", "product_name": "Juice Mango", "category": "Beverage"},
    "manufacturer_and_marketer": {"manufactured_by": None, "country_of_origin": None},
    "dates_and_batch": {"mfg_date": "04/2026", "exp_date": "10/2026", "batch_number": None},
    "pricing_and_quantity": {"mrp": "MRP Rs.65.00 (incl. of all taxes)", "mrp_numerical_inr": 65, "net_quantity": "1 L"},
    "compliance_summary": {},
}


def _completed_inspection() -> int:
    res = client.post("/api/inspections", json={"notes": "report test"}, headers=HEADERS)
    inspection_id = res.json()["id"]

    import uuid
    storage = os.path.join(BASE_DIR, "..", "storage", "test")
    os.makedirs(storage, exist_ok=True)
    path = os.path.join(storage, f"{inspection_id}_{uuid.uuid4().hex[:6]}.jpg")
    import cv2
    import numpy as np
    img = np.full((240, 320, 3), 235, dtype=np.uint8)
    cv2.putText(img, "ACME JUICE MANGO", (30, 100), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    cv2.putText(img, "MRP Rs.65.00", (30, 160), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
    cv2.imwrite(path, img)
    with open(path, "rb") as fh:
        client.post(
            f"/api/inspections/{inspection_id}/images",
            files={"files": ("pack.jpg", fh, "image/jpeg")},
            headers=HEADERS,
        )

    from app.compliance.engine import RuleEngine

    pipeline_result = {
        "per_image_results": [{"full_text": TRANSCRIPT, "text_lines": TRANSCRIPT.splitlines(),
                               "ocr_total_lines": 5, "ocr_latency_sec": 0.4,
                               "token_usage": {}, "source_path": "pack.jpg",
                               "saved_color_path": None, "saved_ocr_path": None}],
        "aggregated_data": AGGREGATED,
        "aggregator_meta": {"model_used": "mock-model"},
        "markdown_report": "# mock",
        "legal_report_markdown": "# mock legal",
        "total_pipeline_time_sec": 1.0,
    }
    pipeline_result["compliance_report"] = RuleEngine().evaluate(pipeline_result).model_dump()

    with patch(
        "app.services.inspection_service._get_pipeline",
        return_value=type("FP", (), {"run_on_images": staticmethod(lambda *a, **k: pipeline_result)}),
    ), patch("app.database.SessionLocal", TestSession):
        client.post(f"/api/inspections/{inspection_id}/process", headers=HEADERS)

    db = TestSession()
    try:
        assert db.get(Inspection, inspection_id).status.value == "COMPLETED"
    finally:
        db.close()
    return inspection_id


class TestReportGeneration:
    def test_pdf_and_docx_from_real_data(self):
        inspection_id = _completed_inspection()

        for fmt, marker in (("pdf", b"%PDF"), ("docx", b"PK")):
            res = client.post(f"/api/reports/generate/{inspection_id}/{fmt}", headers=HEADERS)
            assert res.status_code == 200, res.text
            body = res.json()
            assert body["format"] == fmt
            dl = client.get(body["download_url"], headers=HEADERS)
            assert dl.status_code == 200
            assert dl.content.startswith(marker)

    def test_docx_content_refuses_substitution(self):
        # python-docx readable check: NOT_DETECTED stays "Not detected"
        from docx import Document
        from app.services.report_service import _build_docx

        db = TestSession()
        try:
            inspection = db.get(Inspection, _completed_inspection())
            doc = Document(io.BytesIO(_build_docx(inspection)))
            text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
                cell.text for t in doc.tables for row in t.rows for cell in row.cells
            )
        finally:
            db.close()

        assert "Not detected" in text
        assert "NIRIKSH" in text
        assert "COMPLIANT" in text
        assert "ACME" in text
        assert "no detected value" not in text.lower()

    def test_report_blocked_for_uncompleted(self):
        res = client.post("/api/inspections", json={}, headers=HEADERS)
        draft_id = res.json()["id"]
        res = client.post(f"/api/reports/generate/{draft_id}/pdf", headers=HEADERS)
        assert res.status_code == 409
